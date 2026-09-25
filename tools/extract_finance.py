"""Ekstrak DataPencatatanFinance.xlsx -> services/api/seed-data/finance_*.json (Fase 0.3b).

Sumber: docs/RawData/DataPencatatanFinance.xlsx.
- JU, Kas, Bank: blok saldo awal bank (no. rekening) + seluruh baris jurnal.
- Hutang: daftar vendor + saldo awal/akhir (sheet menulis "Angka hanya contoh" —
  diimpor apa adanya persis dokumen).
- Piutang: daftar customer + saldo awal/akhir.
- Aset: daftar penyusutan fiskal (bangunan + alat berat).
Jalankan dari root repo:  python tools/extract_finance.py
"""
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "RawData" / "DataPencatatanFinance.xlsx"
OUTDIR = ROOT / "services" / "api" / "seed-data"


def s(v):
    return "" if v is None else str(v).strip()


def fnum(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return 0


def iso(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    # ---- JU, Kas, Bank ----
    ws = wb["JU, Kas, Bank"]
    rows = list(ws.iter_rows(values_only=True))
    banks, jurnal = [], []
    for r in rows:
        uraian = s(r[3])
        if uraian.startswith("Saldo Awal Bank"):
            banks.append({
                "nama": uraian,
                "debit": fnum(r[6]), "saldo": fnum(r[8]),
            })
        elif s(r[0]) or s(r[1]) or s(r[2]) or (uraian and not uraian.startswith("BANK") and "Bulan" not in uraian):
            if uraian in ("TANGGAL", "") and not s(r[0]):
                continue
            jurnal.append({
                "tanggal": iso(r[0]), "pihak": s(r[1]), "dok": s(r[2]),
                "uraian": uraian, "akunDb": s(r[4]), "akunKr": s(r[5]),
                "debit": fnum(r[6]), "kredit": fnum(r[7]), "saldo": fnum(r[8]),
            })

    # ---- Hutang / Piutang ----
    def duo(name):
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))[4:]
        out = []
        for r in rows:
            if not s(r[1]):
                continue
            out.append({
                "no": s(r[0]), "nama": s(r[1]),
                "saldoAwal": fnum(r[2]), "saldoAkhir": fnum(r[3]),
            })
        return out

    hutang = duo("Hutang")
    piutang = duo("Piutang")

    # ---- Aset ----
    ws = wb["Aset"]
    rows = list(ws.iter_rows(values_only=True))[6:]  # lewati 6 baris header
    aset = []
    for r in rows:
        if not s(r[1]) and not s(r[0]):
            continue
        aset.append({
            "no": s(r[0]), "nama": s(r[1]), "kelompok": s(r[2]),
            "bln": s(r[3]), "thn": s(r[4]), "nilai": fnum(r[5]),
            "sisaAwal": fnum(r[6]), "metKomersial": s(r[7]),
            "metFiskal": s(r[9]), "susutTahun": fnum(r[10]),
            "akumulasi": s(r[11]),
        })

    OUTDIR.mkdir(parents=True, exist_ok=True)
    (OUTDIR / "finance_bank.json").write_text(
        json.dumps({"banks": banks, "jurnal": jurnal}, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUTDIR / "finance_hutang.json").write_text(
        json.dumps(hutang, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUTDIR / "finance_piutang.json").write_text(
        json.dumps(piutang, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUTDIR / "finance_aset.json").write_text(
        json.dumps(aset, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"bank={len(banks)} jurnal={len(jurnal)} | hutang={len(hutang)} "
          f"total={sum(x['saldoAkhir'] for x in hutang):,.0f} | piutang={len(piutang)} "
          f"total={sum(x['saldoAkhir'] for x in piutang):,.0f} | aset={len(aset)}")


if __name__ == "__main__":
    main()
