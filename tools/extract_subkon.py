"""Ekstrak INVOICE SUBKONTRAKTOR.xlsx -> services/api/seed-data/subkon.json (Fase 0.3a).

Sumber: docs/RawData/Invoice/INVOICE SUBKONTRAKTOR.xlsx (sheet PAK YUSUF).
Isi riil hanya R1-R23 (judul, 3 item harga @100rb, SUB TOTAL 300000,
PPh 0,5% = -1500, TOTAL BAYAR 298500, note PPh variabel).
R108 = formula '=#REF!+#REF!' yatim (rusak di file asli) -> dicatat, tidak diimpor.
Verifikasi: subtotal 300000, pph -1500, total 298500.
Jalankan dari root repo:  python tools/extract_subkon.py
"""
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "RawData" / "Invoice" / "INVOICE SUBKONTRAKTOR.xlsx"
OUT = ROOT / "services" / "api" / "seed-data" / "subkon.json"


def s(v):
    return "" if v is None else str(v).strip()


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["PAK YUSUF"]

    title = s(ws.cell(row=1, column=1).value)
    assert "BG RMN 3324" in title, title

    # Baris item berdimensi/harga: R7, R9-R10, R12, R14 (R11/R13/R15 sub-ket tanpa angka).
    def row(i, n=12):
        return [ws.cell(row=i, column=c).value for c in range(1, n + 1)]

    lines = []
    for i in (7, 9, 10, 12, 14):
        r = row(i)
        lines.append({
            "row": i,
            "uraian": " ".join(s(c) for c in r[:3] if s(c)),
            "cols": [s(c) for c in r],
        })

    subtotal = next((ws.cell(row=17, column=c).value for c in range(1, 22) if isinstance(ws.cell(row=17, column=c).value, (int, float))), None)
    pph = next((ws.cell(row=18, column=c).value for c in range(1, 22) if isinstance(ws.cell(row=18, column=c).value, (int, float))), None)
    total = next((ws.cell(row=19, column=c).value for c in range(1, 22) if isinstance(ws.cell(row=19, column=c).value, (int, float))), None)
    assert int(subtotal) == 300000, subtotal
    assert int(pph) == -1500, pph
    assert int(total) == 298500, total

    note = s(ws.cell(row=23, column=2).value)

    # R108: formula rusak di file asli.
    wb2 = openpyxl.load_workbook(SRC, read_only=True, data_only=False)
    ref_formula = str(wb2["PAK YUSUF"].cell(row=108, column=1).value or "")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({
        "title": title,
        "sub": "Pak Yusuf",
        "vessel": "BG RMN 3324",
        "lines": lines,
        "subtotal": int(subtotal),
        "pph": int(pph),
        "total": int(total),
        "note": note,
        "brokenRef": {"row": 108, "formula": ref_formula},
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"OK subtotal={subtotal} pph={pph} total={total} ref={ref_formula!r} -> {OUT}")


if __name__ == "__main__":
    main()
